"""Import a manager hierarchy and one store login per location from an Excel workbook."""

import argparse
import re
import unicodedata
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import select

from app.core import hash_password
from app.database import SessionLocal
from app.models import (
    Consultant,
    Coordinator,
    Regional,
    Role,
    Store,
    User,
    UserCoordinator,
    UserStore,
)


def slugify(value: str) -> str:
    normalized = unicodedata.normalize("NFKD", value).encode("ascii", "ignore").decode()
    return re.sub(r"-+", "-", re.sub(r"[^a-z0-9]+", "-", normalized.lower())).strip("-")


def store_login(store_name: str) -> str:
    """Create the short username used by a store, e.g. ``maua``."""
    without_code = re.sub(r"^\s*\d+\s*-\s*", "", store_name)
    without_suffix = re.sub(r"\s*-\s*comm\s*$", "", without_code, flags=re.IGNORECASE)
    login = slugify(without_suffix).replace("-", "")
    if not login:
        raise ValueError(f"Não foi possível gerar um login para a loja: {store_name}")
    return login


def get_or_create(session, model, **criteria):
    item = session.scalar(select(model).filter_by(**criteria))
    if item is None:
        item = model(**criteria)
        session.add(item)
        session.flush()
    return item


def is_header(values: tuple[str, str, str, str]) -> bool:
    first, second, third, _ = (value.lower() for value in values)
    return ("nome" in first or "consult" in first) and ("filial" in second or "loja" in second) and "coord" in third


def rows_from_workbook(path: Path) -> tuple[list[tuple[str, str, str]], int]:
    sheet = load_workbook(path, read_only=True, data_only=True).active
    rows: list[tuple[str, str, str]] = []
    seen: set[tuple[str, str, str]] = set()
    duplicates = 0
    for row_number, values in enumerate(sheet.iter_rows(values_only=True), start=1):
        raw = tuple(str(value or "").strip() for value in values[:4])
        raw = raw + ("",) * (4 - len(raw))
        if row_number == 1 and is_header(raw):
            continue
        consultant, store, coordinator, _manager = raw
        if not all((consultant, store, coordinator)):
            raise ValueError(f"Linha {row_number} precisa ter consultor, loja e coordenador.")
        row = (consultant, store, coordinator)
        if row in seen:
            duplicates += 1
            continue
        seen.add(row)
        rows.append(row)
    if not rows:
        raise ValueError("A planilha não possui linhas para importar.")
    return rows, duplicates


def validate_login_collisions(rows: list[tuple[str, str, str]]) -> None:
    logins: dict[str, tuple[str, str]] = {}
    for _, store_name, coordinator_name in rows:
        key = (coordinator_name, store_name)
        login = store_login(store_name)
        other = logins.setdefault(login, key)
        if other != key:
            raise ValueError(f"As lojas {other[1]} e {store_name} geram o mesmo login: {login}")


def assign_manager(session, login: str, name: str, coordinator_ids: set[int]) -> bool:
    manager = session.scalar(select(User).where(User.email == login))
    created = manager is None
    if manager is None:
        manager = User(name=name, email=login, password_hash=hash_password(f"{login}@321"), role=Role.MANAGER, active=True)
        session.add(manager)
        session.flush()
    elif manager.role != Role.MANAGER:
        raise ValueError(f"O login {login} já pertence a um usuário que não é gerente.")
    existing = set(session.scalars(select(UserCoordinator.coordinator_id).where(UserCoordinator.user_id == manager.id)))
    for coordinator_id in coordinator_ids - existing:
        session.add(UserCoordinator(user_id=manager.id, coordinator_id=coordinator_id))
    return created


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--regional", default="SPL")
    parser.add_argument("--manager-login")
    parser.add_argument("--manager-name")
    args = parser.parse_args()

    rows, duplicate_rows = rows_from_workbook(args.workbook)
    validate_login_collisions(rows)
    with SessionLocal() as session:
        try:
            regional = get_or_create(session, Regional, name=args.regional)
            coordinator_by_name: dict[str, Coordinator] = {}
            store_by_key: dict[tuple[str, str], Store] = {}
            for consultant_name, store_name, coordinator_name in rows:
                coordinator = coordinator_by_name.get(coordinator_name)
                if coordinator is None:
                    coordinator = get_or_create(session, Coordinator, regional_id=regional.id, name=coordinator_name)
                    coordinator_by_name[coordinator_name] = coordinator
                store = store_by_key.get((coordinator_name, store_name))
                if store is None:
                    store = get_or_create(session, Store, coordinator_id=coordinator.id, name=store_name)
                    store_by_key[(coordinator_name, store_name)] = store
                get_or_create(session, Consultant, store_id=store.id, name=consultant_name)

            created_users = 0
            for store in store_by_key.values():
                login = store_login(store.name)
                user = session.scalar(select(User).where(User.email == login))
                if user is None:
                    user = User(name=f"Loja {store.name}", email=login, password_hash=hash_password(f"{login}@321"), role=Role.OPERATOR, active=True)
                    session.add(user)
                    session.flush()
                    created_users += 1
                assigned = set(session.scalars(select(UserStore.store_id).where(UserStore.user_id == user.id)))
                if user.role != Role.OPERATOR or assigned not in (set(), {store.id}):
                    raise ValueError(f"O login {login} já está vinculado a outro usuário ou loja.")
                if not assigned:
                    session.add(UserStore(user_id=user.id, store_id=store.id))

            manager_created = False
            if args.manager_login:
                manager_created = assign_manager(session, args.manager_login, args.manager_name or args.manager_login.title(), {item.id for item in coordinator_by_name.values()})
            session.commit()
        except Exception:
            session.rollback()
            raise

    print(f"Importação concluída: {len(rows)} consultores, {len(store_by_key)} lojas, {len(coordinator_by_name)} coordenadores e {created_users} novos logins de loja. Duplicidades ignoradas: {duplicate_rows}. Gerente criado: {manager_created}.")


if __name__ == "__main__":
    main()
